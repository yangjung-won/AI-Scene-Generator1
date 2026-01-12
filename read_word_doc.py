"""
临时脚本：读取Word文档内容
"""
import sys
import os

doc_path = r"C:\Users\Lenovo\Desktop\实验与评估.docx"

try:
    from docx import Document
    doc = Document(doc_path)
    
    print("=" * 80)
    print("Word文档内容：")
    print("=" * 80)
    print()
    
    for i, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip()
        if text:
            print(text)
            print()
    
    # 读取表格
    if doc.tables:
        print("\n" + "=" * 80)
        print("表格内容：")
        print("=" * 80)
        for table_idx, table in enumerate(doc.tables):
            print(f"\n表格 {table_idx + 1}:")
            for row in table.rows:
                row_data = [cell.text.strip() for cell in row.cells]
                print(" | ".join(row_data))
    
    print("\n" + "=" * 80)
    print("文档读取完成")
    print("=" * 80)
    
except ImportError:
    print("错误：需要安装 python-docx 库")
    print("请运行: pip install python-docx")
    sys.exit(1)
except Exception as e:
    print(f"读取文档时出错: {e}")
    import traceback
    traceback.print_exc()
    sys.exit(1)









